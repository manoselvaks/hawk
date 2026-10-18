"""Rebuild the market fixtures as trimmed excerpts of the real Bank of England files.

Usage::

    python tests/market/fixtures/build_fixtures.py <download dir>

where ``<download dir>`` holds the four files as downloaded from the Bank:
``Bank-Rate.asp`` (saved page), ``mpcvoting.xlsx``, ``glcnominalddata.zip`` and
``latest-yield-curve-data.zip`` (URLs in ``hawk.market.sources``).

Every kept cell is the Bank's own value in its original layout; only rows are
dropped. The Bank Rate table is cut into short runs of consecutive changes around
the dates the tests care about, and the voting workbook keeps one decision row per
kept change plus a few holds, so the excerpts stay consistent with each other.
Where a run starts, the "previous rate" is the end of the run before it, not the
true previous rate; the tests only assert on rows inside a run.
"""

from __future__ import annotations

import io
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook

HERE = Path(__file__).parent
BANK_RATE_URL = "https://www.bankofengland.co.uk/boeapps/database/Bank-Rate.asp"

# Runs of consecutive Bank Rate changes (the first row of each run gives the rate before).
BANK_RATE_ROWS = {
    "30 Oct 96", "06 May 97", "06 Jun 97",  # last pre-MPC change, then the first MPC hike
    "06 Nov 97", "04 Jun 98",
    "02 Aug 01", "18 Sep 01",  # unscheduled cut after 11 September
    "10 Apr 08", "08 Oct 08", "06 Nov 08",  # coordinated cut, then -150bp
    "02 Aug 18", "11 Mar 20", "19 Mar 20",  # Covid emergency cuts to 0.1%
    "22 Jun 23", "03 Aug 23",  # the 5.25% peak
}  # fmt: skip
# Decisions kept from the voting sheet: every kept change from June 1997 on, plus holds.
HOLDS = {
    date(1998, 1, 8),  # three dissents recorded only as "Increase"
    date(2020, 3, 26),
    date(2023, 9, 21),  # 5-4 hold at the peak
}
# Yield curve rows kept, per workbook in the archive.
GLC_ROWS = {
    "GLC Nominal daily data_1995 to 1999.xlsx": [
        datetime(1996, 12, 31),  # before the sample: dropped by the collector
        datetime(1997, 1, 2),
        datetime(1997, 1, 3),
    ],
    "GLC Nominal daily data_2016 to 2024.xlsx": [
        datetime(2022, 9, 22),
        datetime(2022, 9, 23),
        datetime(2022, 9, 26),
        datetime(2022, 9, 27),  # the mini-budget spike
    ],
}
GLC_LATEST_ROWS = 2  # first business days of the current month
GLC_HEADER_ROWS = 5  # title, blank, "Maturity", the "years:" row, blank


def trim_bank_rate(page: bytes) -> bytes:
    soup = BeautifulSoup(page, "lxml")
    table = soup.find("table", id="stats-table")
    assert table is not None
    for tr in table.find_all("tr"):
        cells = tr.find_all("td")
        if cells and cells[0].get_text(strip=True) not in BANK_RATE_ROWS:
            tr.decompose()
    kept = len(table.find_all("td")) // 2
    assert kept == len(BANK_RATE_ROWS), f"kept {kept} Bank Rate rows"
    return f"<!-- Excerpt of {BANK_RATE_URL} -->\n<html><body>\n{table}\n</body></html>\n".encode()


def trim_mpc_voting(workbook: bytes) -> bytes:
    """Copy the header rows and the chosen decision rows into a fresh workbook."""
    changes = {datetime.strptime(d, "%d %b %y").date() for d in BANK_RATE_ROWS}
    keep = {d for d in changes if d >= date(1997, 6, 6)} | HOLDS
    source = load_workbook(io.BytesIO(workbook), read_only=True, data_only=True)
    rows = list(source["Bank Rate Decisions"].iter_rows(values_only=True))
    first_decision = next(i for i, row in enumerate(rows) if isinstance(row[1], datetime))
    by_date = {r[1].date(): r for r in rows[first_decision:] if isinstance(r[1], datetime)}
    assert keep <= set(by_date), sorted(keep - set(by_date))
    chosen = [row for day, row in by_date.items() if day in keep]

    book = Workbook()
    sheet = book.worksheets[0]
    sheet.title = "Bank Rate Decisions"
    for row in [*rows[:first_decision], *chosen]:
        sheet.append(list(row))
    return _save(book)


def trim_glc(workbook: bytes, keep: list[datetime] | int) -> bytes:
    """Copy the spot curve sheet's header rows and the chosen dates into a new workbook."""
    source = load_workbook(io.BytesIO(workbook), read_only=True, data_only=True)
    spot = next(name for name in source.sheetnames if name.strip().endswith("spot curve"))
    rows = list(source[spot].iter_rows(values_only=True))
    body = [row for row in rows[GLC_HEADER_ROWS:] if isinstance(row[0], datetime)]
    chosen = body[:keep] if isinstance(keep, int) else [r for r in body if r[0] in keep]
    assert isinstance(keep, int) or len(chosen) == len(keep), spot

    book = Workbook()
    book.worksheets[0].title = "3. spot, short end"  # present in the real files; must be skipped
    sheet = book.create_sheet(spot)
    for row in [*rows[:GLC_HEADER_ROWS], *chosen]:
        sheet.append(list(row))
    return _save(book)


def _zip(files: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def _save(book: Workbook) -> bytes:
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def main(downloads: Path) -> None:
    page = (downloads / "Bank-Rate.asp").read_bytes()
    (HERE / "bank_rate.html").write_bytes(trim_bank_rate(page))
    voting = (downloads / "mpcvoting.xlsx").read_bytes()
    (HERE / "mpcvoting_sample.xlsx").write_bytes(trim_mpc_voting(voting))

    with zipfile.ZipFile(downloads / "glcnominalddata.zip") as archive:
        books = {name: trim_glc(archive.read(name), keep) for name, keep in GLC_ROWS.items()}
    books["GLC Real daily data_2016 to 2024.xlsx"] = b"not a nominal curve: ignored"
    (HERE / "glc_archive_sample.zip").write_bytes(_zip(books))

    latest = "GLC Nominal daily data current month.xlsx"
    with zipfile.ZipFile(downloads / "latest-yield-curve-data.zip") as archive:
        current = trim_glc(archive.read(latest), GLC_LATEST_ROWS)
    (HERE / "glc_latest_sample.zip").write_bytes(_zip({latest: current}))


if __name__ == "__main__":
    main(Path(sys.argv[1]))
